#!/usr/bin/env python3
"""
Audits the generated workbook for the two promises made to the user:

A. LOCKING  - exactly the intended input cells are unlocked, no formula cell is unlocked,
              every sheet is protected, the engine sheet is hidden, the workbook structure
              is locked, and data validation guards every input.
B. IMMUTABILITY - re-writing a later interest rate (value or effective date) cannot change
              any quarter that ended before the revision takes effect.

Usage: python3 audit_workbook.py <template.xlsx> <example.xlsx>
"""
from __future__ import annotations

import shutil
import sys

import formulas
import openpyxl

from build_loan_workbook import (
    CELL_R2_DATE, CELL_R2_PCT, ENTRY_ROWS, QUARTERS, SHEET_ENG, SHEET_HELP, SHEET_MAIN,
    SHEET_SUM, ROW_P1, ROW_P_LAST, qtr_rows,
)

fails: list[str] = []


def check(cond: bool, label: str):
    print(f"  [{'PASS' if cond else 'FAIL'}] {label}")
    if not cond:
        fails.append(label)


# ------------------------------------------------------------------ A. locking audit
def expected_unlocked() -> set[str]:
    cells = set()
    for r in range(ROW_P1, ROW_P_LAST):            # account particulars C:E
        cells |= {f"C{r}", f"D{r}", f"E{r}"}
    for r in range(ROW_P1, ROW_P_LAST):            # basis / lock control H:I
        cells |= {f"H{r}", f"I{r}"}
    cells |= {"K8", "K9", "L7", "L8", "L9"}        # rate register
    for q in range(1, QUARTERS + 1):
        _, bf, e1, e2, ir = qtr_rows(q)
        for r in range(e1, e2 + 1):
            cells |= {f"B{r}", f"C{r}", f"D{r}", f"F{r}"}
        for r in range(bf, ir + 1):
            cells.add(f"L{r}")
    return cells


def audit_locking(path: str):
    wb = openpyxl.load_workbook(path)
    print("\nA. LOCKING AUDIT")
    check(wb.sheetnames == [SHEET_MAIN, SHEET_SUM, SHEET_HELP, SHEET_ENG],
          f"sheet order {wb.sheetnames}")
    for name in wb.sheetnames:
        ws = wb[name]
        check(bool(ws.protection.sheet), f"sheet '{name}' is protected")
        check(bool(ws.protection.password), f"sheet '{name}' protection has a password")
    check(wb[SHEET_ENG].sheet_state == "veryHidden", "calculation engine sheet is veryHidden")
    check(bool(wb.security and wb.security.lockStructure),
          "workbook structure is locked (sheets cannot be added/removed/unhidden)")

    ws = wb[SHEET_MAIN]
    for attr in ("insertRows", "deleteRows", "insertColumns", "deleteColumns", "sort",
                 "formatCells"):
        check(getattr(ws.protection, attr) is True, f"ledger: '{attr}' is blocked")

    unlocked, formula_unlocked = set(), []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=20):
        for c in row:
            if c.protection and c.protection.locked is False:
                unlocked.add(c.coordinate)
                if isinstance(c.value, str) and c.value.startswith("="):
                    formula_unlocked.append(c.coordinate)
    exp = expected_unlocked()
    check(unlocked == exp,
          f"exactly the intended {len(exp)} cells are typeable "
          f"(found {len(unlocked)}; unexpected={sorted(unlocked - exp)[:6]}, "
          f"missing={sorted(exp - unlocked)[:6]})")
    check(not formula_unlocked, f"no formula cell is typeable ({formula_unlocked[:5]})")

    # every input cell is covered by a data validation rule
    covered = set()
    for d in ws.data_validations.dataValidation:
        for rng in d.sqref.ranges:
            got = ws[str(rng)]
            if isinstance(got, openpyxl.cell.cell.Cell):
                covered.add(got.coordinate)
                continue
            for row in got:
                row = row if isinstance(row, tuple) else (row,)
                for c in row:
                    covered.add(c.coordinate)
    guard_needed = {c for c in exp if c[0] in "BFKLH"}
    missing = sorted(guard_needed - covered)
    check(not missing, f"every date/amount/rate input has a validation rule "
                       f"(unguarded: {missing[:8]})")

    # formulas must not be hard-coded values in the ledger's computed columns
    # Column K on a quarter-end row is intentionally the constant 0: accrued interest is
    # reset once it has been capitalised into the balance.
    hardcoded = []
    for q in range(1, QUARTERS + 1):
        _, bf, e1, e2, ir = qtr_rows(q)
        for r in range(e1, e2 + 1):
            for col in ("G", "I", "J", "K", "M", "N"):
                v = ws[f"{col}{r}"].value
                if not (isinstance(v, str) and v.startswith("=")):
                    hardcoded.append(f"{col}{r}")
        for col in ("E", "G", "I", "J", "M", "N"):
            v = ws[f"{col}{ir}"].value
            if not (isinstance(v, str) and v.startswith("=")):
                hardcoded.append(f"{col}{ir}")
        check_bf = ws[f"G{bf}"].value
        if not (isinstance(check_bf, str) and check_bf.startswith("=")):
            hardcoded.append(f"G{bf}")
        if ws[f"K{ir}"].value != 0:
            hardcoded.append(f"K{ir} should reset accrued interest to 0")
    check(not hardcoded, f"balance/day-count/interest columns are all live formulas "
                         f"({hardcoded[:5]})")
    check(len(exp) == 15 + 10 + 5 + QUARTERS * (ENTRY_ROWS * 4 + ENTRY_ROWS + 2),
          "input-cell count matches the documented layout")


# ------------------------------------------------------- B. immutability of past quarters
def quarter_figures(path: str) -> dict[int, tuple[float, float]]:
    xl = formulas.ExcelModel().loads(path).finish()
    sol = xl.calculate()
    book = path.split("/")[-1].upper()
    want = {}
    for q in range(1, QUARTERS + 1):
        _, _, _, _, ir = qtr_rows(q)
        want[f"E{ir}"] = q
        want[f"G{ir}"] = q
    out: dict[int, dict[str, float]] = {q: {} for q in range(1, QUARTERS + 1)}
    for key, val in sol.items():
        k = key.upper()
        if f"[{book}]" not in k or SHEET_MAIN.upper() not in k:
            continue
        ref = k.split("!")[-1].replace("'", "")
        if ref in want:
            try:
                out[want[ref]][ref[0]] = float(val.value[0, 0])
            except Exception:
                pass
    return {q: (v.get("E"), v.get("G")) for q, v in out.items()}


def audit_immutability(example: str):
    print("\nB. IMMUTABILITY AUDIT (recalculated from the real workbook formulas)")
    base = quarter_figures(example)

    scenarios = [
        ("2nd rate raised 10.50% -> 25.00%", {CELL_R2_PCT: 25.0}, 3),
        ("2nd rate effective date pushed 01-07-2025 -> 01-09-2025",
         {CELL_R2_DATE: __import__("datetime").date(2025, 9, 1)}, 3),
    ]
    for label, patch, unaffected_upto in scenarios:
        tampered = example.replace(".xlsx", "._tamper.xlsx")
        shutil.copy(example, tampered)
        wb = openpyxl.load_workbook(tampered)
        for cell, value in patch.items():
            wb[SHEET_MAIN][cell] = value
        wb.save(tampered)
        after = quarter_figures(tampered)

        print(f"\n  scenario: {label}")
        frozen_ok, later_changed = True, False
        for q in range(1, QUARTERS + 1):
            same = (abs(base[q][0] - after[q][0]) < 0.005
                    and abs(base[q][1] - after[q][1]) < 0.005)
            if q <= unaffected_upto:
                print(f"    Q{q} (ended before the revision): interest "
                      f"{base[q][0]:>12,.2f} -> {after[q][0]:>12,.2f}  "
                      f"{'unchanged' if same else 'CHANGED'}")
                frozen_ok &= same
            elif not same:
                later_changed = True
        check(frozen_ok, f"quarters closed before the revision are untouched ({label})")
        check(later_changed, f"quarters after the revision do pick up the new rate ({label})")
        import os
        os.remove(tampered)


if __name__ == "__main__":
    audit_locking(sys.argv[1])
    audit_immutability(sys.argv[2])
    print(f"\n{len(fails)} failure(s)")
    for f in fails:
        print("  FAIL:", f)
    sys.exit(1 if fails else 0)
